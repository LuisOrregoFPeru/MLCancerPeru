"""
Dashboard "Cáncer en el Tiempo — Perú (INEN)"
==============================================
Inspirado en el visor "Cancer Over Time" de GLOBOCAN/IARC, construido
sobre los datos del INEN (Instituto Nacional de Enfermedades
Neoplásicas) para los 25 departamentos del Perú + total nacional.

Ejecutar localmente:
    streamlit run app.py

Estructura del proyecto:
    app.py                -> esta app (interfaz)
    data_processing.py    -> carga, limpieza y actualización de datos
    data/                 -> Excel original + caché limpia
    requirements.txt      -> dependencias
    .streamlit/config.toml-> tema visual
"""

from __future__ import annotations

import colorsys
import io

import numpy as np
import pandas as pd
import plotly.graph_objects as go
import streamlit as st

from data_processing import (
    ALL_SITES_LABEL,
    PERU_LABEL,
    append_new_data,
    departamentos,
    load_data,
    localizaciones,
)
from stats_analysis import (
    chow_test_arbitrary_break,
    detect_breakpoint,
    linear_trend,
    mann_kendall_trend,
    project_series,
    smooth_series,
)

# Generación de GIF para el ranking animado: dependencias opcionales — si no
# están disponibles, la animación interactiva de Plotly sigue funcionando y
# solo se oculta el botón de descarga de GIF.
try:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import imageio.v2 as imageio

    _GIF_EXPORT_AVAILABLE = True
except ImportError:  # pragma: no cover
    _GIF_EXPORT_AVAILABLE = False

# Generación del Reporte Ejecutivo (PDF / PPTX): dependencias opcionales,
# puras en Python (sin binarios de sistema como Chrome/LibreOffice) para
# no repetir problemas de despliegue en Streamlit Cloud.
from datetime import datetime

try:
    from reportlab.lib import colors as rl_colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        Image as RLImage,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    _PDF_EXPORT_AVAILABLE = True
except ImportError:  # pragma: no cover
    _PDF_EXPORT_AVAILABLE = False

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    _PPTX_EXPORT_AVAILABLE = True
except ImportError:  # pragma: no cover
    _PPTX_EXPORT_AVAILABLE = False

# ---------------------------------------------------------------------------
# Configuración general de la página
# ---------------------------------------------------------------------------

st.set_page_config(
    page_title="Cáncer en el Tiempo | INEN Perú",
    page_icon="🎗️",
    layout="wide",
    initial_sidebar_state="expanded",
)

PRIMARY_COLOR = "#8b1a4f"      # línea principal (similar al magenta de GLOBOCAN)
SECONDARY_COLOR = "#2f6fa8"    # línea secundaria (azul)
ACCENT_COLOR = "#1f2d5c"       # navy institucional

CUSTOM_CSS = f"""
<style>
    .main .block-container {{ padding-top: 1.3rem; max-width: 1300px; }}
    h1, h2, h3 {{ color: {ACCENT_COLOR}; }}
    .app-header {{
        background: linear-gradient(90deg, {ACCENT_COLOR} 0%, {PRIMARY_COLOR} 100%);
        padding: 1.1rem 1.6rem;
        border-radius: 10px;
        color: white;
        margin-bottom: 1.2rem;
    }}
    .app-header h1 {{ color: white; margin: 0; font-size: 1.6rem; }}
    .app-header p {{ color: #e8e6f0; margin: 0.2rem 0 0 0; font-size: 0.95rem; }}
    div[data-testid="stMetric"] {{
        background: #f7f7fb;
        border: 1px solid #e6e6ef;
        border-radius: 10px;
        padding: 0.6rem 0.9rem;
    }}
    .source-note {{ font-size: 0.8rem; color: #666; margin-top: 0.4rem; }}
</style>
"""
st.markdown(CUSTOM_CSS, unsafe_allow_html=True)


# ---------------------------------------------------------------------------
# Datos
# ---------------------------------------------------------------------------

@st.cache_data(show_spinner="Cargando datos del INEN...")
def get_data() -> pd.DataFrame:
    return load_data()


df = get_data()

st.markdown(
    f"""
    <div class="app-header">
        <h1>🎗️ Cáncer en el Tiempo — Perú</h1>
        <p>Casos nuevos de cáncer registrados por el INEN ({int(df['Anio'].min())}–{int(df['Anio'].max())})</p>
    </div>
    """,
    unsafe_allow_html=True,
)

DEFAULT_PALETTE = [
    PRIMARY_COLOR, SECONDARY_COLOR, "#e07a2c", "#3ba776", "#8850c4",
    "#c94141", "#4aa8c9", "#c9a13b", "#5c6ac4", "#7d7d7d",
]


def format_region_list(regions: list[str], max_show: int = 3) -> str:
    """Muestra los nombres de las regiones seleccionadas (no solo el
    conteo); si hay muchas, trunca y agrega '+N más' para no romper
    el layout de la tarjeta."""
    if not regions:
        return "—"
    if len(regions) <= max_show:
        return ", ".join(regions)
    return ", ".join(regions[:max_show]) + f" +{len(regions) - max_show} más"


def shades_of(hex_color: str, n: int, l_min: float = 0.30, l_max: float = 0.80) -> list[str]:
    """Genera `n` matices del mismo color (variando solo la luminosidad),
    del más oscuro (primer puesto en el ranking) al más claro. Se usa en
    el gráfico de ranking para que el orden se perciba visualmente sin
    necesitar colores distintos por barra."""
    hex_color = hex_color.lstrip("#")
    r, g, b = (int(hex_color[i : i + 2], 16) / 255 for i in (0, 2, 4))
    h, _l, s = colorsys.rgb_to_hls(r, g, b)
    lights = [l_min] if n <= 1 else np.linspace(l_min, l_max, n)
    colors = []
    for lig in lights:
        rr, gg, bb = colorsys.hls_to_rgb(h, float(lig), s)
        colors.append(
            "#%02x%02x%02x" % (int(round(rr * 255)), int(round(gg * 255)), int(round(bb * 255)))
        )
    return colors


def dept_casos_series(dept: str, site_val: str, year_lo: int, year_hi: int) -> pd.DataFrame:
    """Devuelve las filas Anio/Casos para un departamento y una
    localización del tumor primario, en el rango de años dado.

    Si `dept` es "Perú (Total nacional)" y el usuario configuró
    exclusiones en el panel lateral (variable global `peru_exclude`,
    p. ej. para no contar 'Extranjero'), esos casos se restan del total
    nacional tal como lo reporta el INEN. Si el departamento excluido no
    tiene dato para un año dado, no se resta nada ese año (se asume 0,
    no una ausencia de dato)."""
    base = df[
        df["Localizacion"].eq(site_val)
        & df["Departamento"].eq(dept)
        & df["Anio"].between(year_lo, year_hi)
    ][["Anio", "Casos"]].copy()

    if dept == PERU_LABEL and peru_exclude:
        excl = (
            df[
                df["Localizacion"].eq(site_val)
                & df["Departamento"].isin(peru_exclude)
                & df["Anio"].between(year_lo, year_hi)
            ]
            .groupby("Anio")["Casos"]
            .sum(min_count=1)
        )

        def _subtract(row):
            if pd.isna(row["Casos"]):
                return row["Casos"]
            excl_val = excl.get(row["Anio"], 0)
            # Si el/los departamento(s) excluidos no reportaron dato ese
            # año (NaN, no cero), se asume que no hay nada que restar en
            # vez de propagar el NaN y "borrar" el total nacional de ese año.
            if pd.isna(excl_val):
                excl_val = 0
            return row["Casos"] - excl_val

        base["Casos"] = base.apply(_subtract, axis=1)
    return base


_RANK_MEDALS = ["🥇", "🥈", "🥉"]


def get_year_ranking(
    dept: str,
    year: int,
    exclude: list[str] | None = None,
    top_n: int | None = None,
    peru_excl_override: list[str] | None = None,
) -> pd.DataFrame:
    """Devuelve el ranking (descendente por N° de casos) de
    localizaciones del tumor primario para un departamento y año dados.

    Aplica el ajuste de exclusión del total nacional (cuando `dept` es
    Perú) usando `peru_excl_override` si se especifica, o si no la
    variable global `peru_exclude` configurada en el panel lateral.
    Opcionalmente excluye localizaciones específicas o recorta a las
    primeras `top_n`. Se usa tanto en la pestaña "Ranking por año" como
    en "Ranking animado", para no duplicar la lógica."""
    exclude = exclude or []
    effective_peru_exclude = (
        peru_excl_override if peru_excl_override is not None else peru_exclude
    )
    data = (
        df[
            (df["Anio"] == year)
            & (df["Departamento"] == dept)
            & (df["Localizacion"] != ALL_SITES_LABEL)
            & (~df["Localizacion"].isin(exclude))
        ][["Localizacion", "Casos"]]
        .dropna(subset=["Casos"])
        .copy()
    )

    if dept == PERU_LABEL and effective_peru_exclude:
        excl = (
            df[
                (df["Anio"] == year)
                & (df["Departamento"].isin(effective_peru_exclude))
                & (df["Localizacion"] != ALL_SITES_LABEL)
            ]
            .groupby("Localizacion")["Casos"]
            .sum(min_count=1)
        )

        def _sub(row):
            e = excl.get(row["Localizacion"], 0)
            if pd.isna(e):
                e = 0
            return row["Casos"] - e

        data["Casos"] = data.apply(_sub, axis=1)

    data = data.sort_values("Casos", ascending=False).reset_index(drop=True)
    if top_n is not None:
        data = data.head(top_n)
    return data


def top3_localizaciones(dept: str, year: int) -> str:
    """Texto (HTML, para hover de Plotly) con el top 3 de localizaciones
    del tumor primario con más casos para un departamento y año dados.

    Si `dept` es Perú y hay exclusiones configuradas (variable global
    `peru_exclude`), se aplican de la misma forma que en el resto del
    dashboard antes de calcular el ranking."""
    sub = (
        df[
            (df["Anio"] == year)
            & (df["Departamento"] == dept)
            & (df["Localizacion"] != ALL_SITES_LABEL)
        ][["Localizacion", "Casos"]]
        .dropna(subset=["Casos"])
        .copy()
    )

    if dept == PERU_LABEL and peru_exclude:
        excl = (
            df[
                (df["Anio"] == year)
                & (df["Departamento"].isin(peru_exclude))
                & (df["Localizacion"] != ALL_SITES_LABEL)
            ]
            .groupby("Localizacion")["Casos"]
            .sum(min_count=1)
        )

        def _sub(row):
            e = excl.get(row["Localizacion"], 0)
            if pd.isna(e):
                e = 0
            return row["Casos"] - e

        sub["Casos"] = sub.apply(_sub, axis=1)

    sub = sub.sort_values("Casos", ascending=False).head(3)
    if sub.empty:
        return "Sin datos por localización"
    return "<br>".join(
        f"{_RANK_MEDALS[i]} {row.Localizacion}: {row.Casos:,.0f}"
        for i, row in enumerate(sub.itertuples())
    )

# ---------------------------------------------------------------------------
# Barra lateral — controles (equivalente al panel "Display" de GLOBOCAN)
# ---------------------------------------------------------------------------

with st.sidebar:
    st.header("⚙️ Panel de control")

    with st.expander("📥 Incorporar datos recientes", expanded=False):
        st.caption(
            "Sube un archivo .xlsx o .csv con columnas "
            "`Departamento`, `localizacion` y una columna por cada año "
            "nuevo (ej. `2024`) con el N° de casos. Se combinará con el "
            "histórico; si un dato ya existía, el nuevo lo reemplaza."
        )
        new_file = st.file_uploader(
            "Archivo con datos nuevos", type=["xlsx", "xls", "csv"], key="uploader"
        )
        if new_file is not None:
            tmp_path = f"/tmp/{new_file.name}"
            with open(tmp_path, "wb") as f:
                f.write(new_file.getbuffer())
            try:
                updated = append_new_data(tmp_path, save=True)
                st.success(
                    f"Datos actualizados: {len(updated):,} filas totales. "
                    "Recarga la página para verlos reflejados en los gráficos."
                )
                get_data.clear()
            except Exception as exc:  # noqa: BLE001
                st.error(f"No se pudo procesar el archivo: {exc}")

    st.subheader("Localización del tumor primario")
    site = st.selectbox(
        "Selecciona una localización del tumor primario",
        options=localizaciones(df),
        index=0,
    )

    show_extra_sites = st.checkbox(
        "Comparar varias localizaciones (gráficos adicionales)",
        value=False,
        help=(
            "Agrega un gráfico adicional por cada localización del tumor "
            "primario que selecciones, para evaluarlas simultáneamente sin "
            "perder la selección principal de arriba."
        ),
    )
    extra_sites: list[str] = []
    if show_extra_sites:
        extra_sites = st.multiselect(
            "Localizaciones adicionales a mostrar",
            options=localizaciones(df),
            default=[],
            max_selections=6,
            key="extra_sites",
        )

    st.subheader("Departamento de residencia a comparar")
    depts_available = departamentos(df)
    default_depts = [PERU_LABEL]
    depts_selected = st.multiselect(
        "Selecciona uno o más departamentos de residencia",
        options=depts_available,
        default=default_depts,
        max_selections=10,
    )
    if not depts_selected:
        depts_selected = default_depts

    peru_exclude: list[str] = []
    if PERU_LABEL in depts_selected:
        peru_exclude = st.multiselect(
            "Excluir del total nacional (Perú)",
            options=[d for d in depts_available if d != PERU_LABEL],
            default=[],
            key="peru_exclude",
            help=(
                "Ej.: si no quieres que los casos de 'Extranjero' cuenten "
                "dentro del total de Perú, selecciónalo aquí. Se resta del "
                "total nacional (reportado por el INEN) para la localización "
                "y los años seleccionados, y afecta también al ranking, "
                "Mann-Kendall, quiebres y proyección cuando el departamento "
                "analizado es Perú."
            ),
        )

    st.subheader("Periodo")
    year_min, year_max = int(df["Anio"].min()), int(df["Anio"].max())
    year_range = st.slider(
        "Rango de años",
        min_value=year_min,
        max_value=year_max,
        value=(year_min, year_max),
        step=1,
    )

    st.subheader("Visualización")
    chart_type = st.radio(
        "Tipo de gráfico", ["Línea", "Barras"], horizontal=True
    )
    show_markers = st.checkbox("Mostrar marcadores", value=True)
    show_labels = st.checkbox("Mostrar etiquetas en cada punto", value=False)
    show_rank_labels = st.checkbox(
        "Mostrar etiquetas con ranking",
        value=False,
        help=(
            "Al pasar el cursor sobre un punto, muestra además el top 3 "
            "de localizaciones del tumor primario con más casos ese año "
            "para el departamento correspondiente (🥇🥈🥉)."
        ),
    )
    log_scale = st.checkbox("Escala logarítmica (eje Y)", value=False)
    show_smooth = st.checkbox(
        "Mostrar línea temporal suavizada (LOWESS)",
        value=False,
        help=(
            "Suaviza la serie con regresión local ponderada (LOWESS), "
            "sin imponer una forma lineal, para ver la tendencia general "
            "sin el ruido año a año."
        ),
    )
    smooth_frac = 0.3
    if show_smooth:
        smooth_frac = st.slider(
            "Nivel de suavizado", 0.1, 0.9, 0.3, 0.05,
            help="Valores bajos siguen más de cerca los datos; valores altos generan una curva más suave.",
        )

    st.subheader("Análisis estadístico")
    show_trend = st.checkbox("Línea de tendencia (regresión lineal)", value=False)
    show_breakpoint = st.checkbox(
        "Detectar punto de quiebre (cambio estructural)", value=False,
        help=(
            "Ajusta una regresión segmentada de dos tramos y aplica un "
            "test de Chow para hallar el año en el que la tendencia "
            "cambia de forma estadísticamente significativa."
        ),
    )
    breakpoint_target = None
    if show_breakpoint and depts_selected:
        breakpoint_target = st.selectbox(
            "Departamento a analizar para el quiebre",
            options=depts_selected,
            index=0,
            help="El análisis de quiebre se calcula para una sola serie a la vez.",
        )

    show_mk = st.checkbox(
        "Test de tendencia Mann-Kendall (robusto a autocorrelación)",
        value=False,
        help=(
            "A diferencia de la regresión lineal (que puede mostrar "
            "significancia artificial cuando los residuos están "
            "autocorrelacionados), este test no paramétrico corrige la "
            "varianza por autocorrelación (modificación de Hamed-Rao) "
            "y responde de forma más confiable si existe o no una "
            "tendencia monótona. Incluye la pendiente de Sen, robusta "
            "a valores atípicos."
        ),
    )
    mk_target = None
    mk_year_range = (year_min, year_max)
    if show_mk and depts_selected:
        mk_target = st.selectbox(
            "Departamento a analizar (Mann-Kendall)",
            options=depts_selected,
            index=0,
            key="mk_target_select",
        )
        mk_year_range = st.slider(
            "Rango de años a evaluar (Mann-Kendall)",
            min_value=year_min,
            max_value=year_max,
            value=year_range,
            step=1,
            key="mk_year_range",
            help=(
                "Acota el test a un sub-periodo específico (por ejemplo, "
                "antes o después de un cambio de política), independiente "
                "del rango general del gráfico."
            ),
        )

    show_arbitrary_break = st.checkbox(
        "Evaluar un quiebre en un año específico (evento)",
        value=False,
        help=(
            "A diferencia de la detección automática, aquí indicas tú "
            "el año de un evento (p. ej. una ley aprobada o la "
            "inauguración de un hospital oncológico) y el test de Chow "
            "evalúa si la tendencia cambió realmente a partir de ese "
            "punto. Puedes excluir años de 'implementación' donde el "
            "efecto todavía no se consolida."
        ),
    )
    arb_target = None
    arb_break_year = None
    arb_lag = 0
    if show_arbitrary_break and depts_selected:
        arb_target = st.selectbox(
            "Departamento a analizar (quiebre por evento)",
            options=depts_selected,
            index=0,
            key="arb_target_select",
        )
        arb_break_year = st.number_input(
            "Año del evento",
            min_value=year_min,
            max_value=year_max,
            value=min(max(year_min + 1, (year_min + year_max) // 2), year_max),
            step=1,
            key="arb_break_year",
            help="Ej.: el año en que se aprobó una ley o se inauguró un hospital oncológico.",
        )
        arb_lag = st.slider(
            "Años de implementación a excluir tras el evento",
            0, 5, 0,
            key="arb_lag",
            help=(
                "Si el efecto tarda en consolidarse (p. ej. una ley "
                "aprobada en 2015 que recién se aplica plenamente desde "
                "2017), usa 2 para excluir 2015-2016 del análisis y no "
                "diluir la comparación antes/después."
            ),
        )

    st.subheader("Proyección de casos")
    show_projection = st.checkbox(
        "Proyectar casos a futuro (con análisis de sensibilidad)",
        value=False,
        help=(
            "Muestra 3 escenarios (conservador, recomendado y lineal) "
            "en vez de un único número, para no sobreestimar. El "
            "escenario recomendado usa suavizado exponencial con "
            "tendencia amortiguada (Holt damped trend)."
        ),
    )
    proj_target = None
    proj_horizon = 5
    if show_projection and depts_selected:
        proj_target = st.selectbox(
            "Departamento a proyectar",
            options=depts_selected,
            index=0,
            key="proj_target_select",
        )
        proj_horizon = st.slider("Horizonte de proyección (años)", 1, 10, 5)

    with st.expander("🎨 Personalizar colores"):
        custom_colors = st.checkbox("Elegir color por departamento", value=False)
        color_overrides: dict[str, str] = {}
        if custom_colors:
            for i, d in enumerate(depts_selected):
                color_overrides[d] = st.color_picker(
                    f"Color · {d}", value=DEFAULT_PALETTE[i % len(DEFAULT_PALETTE)]
                )

    st.divider()
    st.caption(
        "Fuente: Instituto Nacional de Enfermedades Neoplásicas (INEN), Perú. "
        "Panel construido con Streamlit + Plotly."
    )

# ---------------------------------------------------------------------------
# Filtrado
# ---------------------------------------------------------------------------

# Se construye departamento por departamento (en vez de un filtro simple)
# para poder aplicar el ajuste de exclusión del total nacional (Perú menos
# los departamentos marcados en `peru_exclude`, p. ej. "Extranjero").
_filtered_frames = []
for _dept in depts_selected:
    _sub = dept_casos_series(_dept, site, year_range[0], year_range[1]).copy()
    _sub["Departamento"] = _dept
    _sub["Localizacion"] = site
    _filtered_frames.append(_sub)

if _filtered_frames:
    filtered = pd.concat(_filtered_frames, ignore_index=True)[
        ["Departamento", "Localizacion", "Anio", "Casos"]
    ].sort_values(["Departamento", "Anio"])
else:
    filtered = df.iloc[0:0][["Departamento", "Localizacion", "Anio", "Casos"]]

# ---------------------------------------------------------------------------
# Resumen general + KPIs
# ---------------------------------------------------------------------------

available_years_kpi = (
    sorted(filtered["Anio"].unique().tolist()) if not filtered.empty else [year_min, year_max]
)
default_year_a = available_years_kpi[0]
default_year_b = available_years_kpi[-1]

# Se leen los años de comparación desde session_state (si el usuario ya
# interactuó con los selectores más abajo) para poder mostrar la tarjeta
# "Casos en <año>" en la primera fila, antes de dibujar los selectores.
year_a = st.session_state.get("kpi_year_a", default_year_a)
year_b = st.session_state.get("kpi_year_b", default_year_b)
if year_a not in available_years_kpi:
    year_a = default_year_a
if year_b not in available_years_kpi:
    year_b = default_year_b

total_year_b = filtered.loc[filtered["Anio"] == year_b, "Casos"].sum(min_count=1)
total_year_a = filtered.loc[filtered["Anio"] == year_a, "Casos"].sum(min_count=1)
delta_pct = (
    ((total_year_b - total_year_a) / total_year_a * 100) if total_year_a else float("nan")
)

row1_col1, row1_col2, row1_col3 = st.columns(3)
row1_col1.metric("Localización del tumor primario", site)
row1_col2.metric("Departamento(s) de residencia", format_region_list(depts_selected))
row1_col3.metric(
    f"Casos en {year_b}",
    f"{int(total_year_b):,}" if pd.notna(total_year_b) else "s/d",
)

st.markdown(
    """
    <div style="background:#1f2d5c; color:white; padding:0.45rem 1rem;
                border-radius:8px; margin:0.9rem 0 0.7rem 0; font-weight:600;
                font-size:0.95rem; letter-spacing:0.03em;">
        📊 KPI
    </div>
    """,
    unsafe_allow_html=True,
)

cmp1, cmp2 = st.columns(2)
with cmp1:
    year_a = st.selectbox(
        "Año base (comparar desde)",
        options=available_years_kpi,
        index=available_years_kpi.index(year_a),
        key="kpi_year_a",
    )
with cmp2:
    year_b = st.selectbox(
        "Año a comparar (comparar hasta)",
        options=available_years_kpi,
        index=available_years_kpi.index(year_b),
        key="kpi_year_b",
    )

# Recalcular con los valores confirmados por los selectores (mismos que
# arriba salvo que el usuario los acabe de cambiar en este rerun)
total_year_b = filtered.loc[filtered["Anio"] == year_b, "Casos"].sum(min_count=1)
total_year_a = filtered.loc[filtered["Anio"] == year_a, "Casos"].sum(min_count=1)
delta_pct = (
    ((total_year_b - total_year_a) / total_year_a * 100) if total_year_a else float("nan")
)
delta_abs = (
    (total_year_b - total_year_a)
    if pd.notna(total_year_a) and pd.notna(total_year_b)
    else float("nan")
)
n_years_cmp = year_b - year_a
if n_years_cmp != 0 and pd.notna(total_year_a) and total_year_a > 0 and pd.notna(total_year_b) and total_year_b > 0:
    cagr = (total_year_b / total_year_a) ** (1 / n_years_cmp) - 1
else:
    cagr = float("nan")

# Serie agregada (suma de todos los departamentos seleccionados) para
# los KPIs que miran todo el periodo filtrado, no solo los 2 años
# comparados: año pico y tendencia resumida.
year_series = filtered.groupby("Anio")["Casos"].sum(min_count=1).dropna()
if not year_series.empty:
    peak_year = int(year_series.idxmax())
    peak_val = float(year_series.max())
else:
    peak_year, peak_val = None, float("nan")

mk_kpi = mann_kendall_trend(year_series.index.to_numpy(), year_series.to_numpy())

kcol1, kcol2, kcol3, kcol4, kcol5 = st.columns(5)
kcol1.metric(
    f"Variación {year_a} → {year_b}",
    f"{delta_pct:+.1f}%" if pd.notna(delta_pct) else "s/d",
)
kcol2.metric(
    "Cambio absoluto",
    f"{delta_abs:+,.0f} casos" if pd.notna(delta_abs) else "s/d",
)
kcol3.metric(
    "CAGR (crec. anual compuesto)",
    f"{cagr * 100:+.1f}%/año" if pd.notna(cagr) else "s/d",
)
kcol4.metric(
    "Año pico",
    f"{peak_year} · {peak_val:,.0f}" if peak_year is not None else "s/d",
)
if mk_kpi is not None:
    trend_icon = {
        "increasing": "📈 Creciente",
        "decreasing": "📉 Decreciente",
        "no trend": "➖ Sin tendencia",
    }.get(mk_kpi.trend, mk_kpi.trend)
    kcol5.metric(
        "Tendencia (Mann-Kendall)",
        trend_icon,
        delta=f"p = {mk_kpi.p_value:.3f}",
        delta_color="off",
    )
else:
    kcol5.metric("Tendencia (Mann-Kendall)", "s/d")

# ---------------------------------------------------------------------------
# Gráfico principal (equivalente al "Graphic" tab de GLOBOCAN)
# ---------------------------------------------------------------------------

(
    tab_graph, tab_table, tab_ranking, tab_ranking_anim, tab_projection,
    tab_downloads, tab_report,
) = st.tabs(
    [
        "📈 Gráfico", "📋 Tabla", "🏆 Ranking por año", "🎬 Ranking animado",
        "🔮 Proyección", "⬇️ Descargas", "📑 Reporte ejecutivo",
    ]
)


def render_site_analysis_chart(
    site_val: str,
    filtered_local: pd.DataFrame,
    compact: bool = False,
    key_suffix: str = "",
) -> None:
    """Construye y renderiza el gráfico de una localización del tumor
    primario, con línea de tendencia, suavizado LOWESS, quiebre
    automático, Mann-Kendall y quiebre por evento —  todo según los
    controles del panel lateral. La usan tanto el gráfico principal
    como los gráficos adicionales de otras localizaciones, así todos
    comparten exactamente el mismo análisis estadístico y la misma
    personalización de colores.

    `compact=True` reduce tamaños de fuente/altura y agrupa los
    resultados estadísticos en un expander, para el layout en columnas
    de los gráficos adicionales.
    """
    if filtered_local.empty or filtered_local["Casos"].dropna().empty:
        st.info(f"No hay datos para **{site_val}** con la selección actual.")
        return

    fig = go.Figure()
    for i, dept in enumerate(depts_selected):
        sub = filtered_local[filtered_local["Departamento"] == dept]
        color = (
            color_overrides.get(dept, DEFAULT_PALETTE[i % len(DEFAULT_PALETTE)])
            if custom_colors
            else DEFAULT_PALETTE[i % len(DEFAULT_PALETTE)]
        )
        mode = "lines+markers" if show_markers else "lines"
        if show_labels:
            mode += "+text"
        raw_opacity = 0.35 if show_smooth else 1.0

        # "Mostrar etiquetas con ranking" es la única opción que NO se
        # aplica a los gráficos compactos de "Comparar con otras
        # localizaciones" — solo al gráfico principal.
        if show_rank_labels and not compact:
            rank_customdata = [top3_localizaciones(dept, y) for y in sub["Anio"]]
            hover_tpl = (
                "%{x}: %{y:,.0f} casos<br><br><b>Top 3 en " + dept + "</b><br>"
                "%{customdata}<extra></extra>"
            )
        else:
            rank_customdata = None
            hover_tpl = "%{x}: %{y:,.0f} casos<extra>" + dept + "</extra>"

        if chart_type == "Línea":
            fig.add_trace(
                go.Scatter(
                    x=sub["Anio"],
                    y=sub["Casos"],
                    name=dept,
                    mode=mode,
                    line=dict(width=2.0 if compact else 2.5, color=color),
                    marker=dict(size=4 if compact else 5),
                    opacity=raw_opacity,
                    text=sub["Casos"].map(lambda v: f"{v:,.0f}" if pd.notna(v) else ""),
                    textposition="top center",
                    textfont=dict(size=8 if compact else 10, color=color),
                    connectgaps=True,
                    customdata=rank_customdata,
                    hovertemplate=hover_tpl,
                )
            )
        else:
            fig.add_trace(
                go.Bar(
                    x=sub["Anio"],
                    y=sub["Casos"],
                    name=dept,
                    marker_color=color,
                    opacity=raw_opacity,
                    text=sub["Casos"].map(lambda v: f"{v:,.0f}" if pd.notna(v) else "")
                    if show_labels
                    else None,
                    textposition="outside",
                    customdata=rank_customdata,
                    hovertemplate=hover_tpl,
                )
            )

        if show_trend:
            trend = linear_trend(sub["Anio"].to_numpy(), sub["Casos"].to_numpy())
            if trend is not None:
                fig.add_trace(
                    go.Scatter(
                        x=trend.x,
                        y=trend.y_pred,
                        name=f"Tendencia · {dept} (R²={trend.r2:.2f})",
                        mode="lines",
                        line=dict(width=1.6, color=color, dash="dot"),
                        hoverinfo="skip",
                    )
                )

        if show_smooth:
            sm = smooth_series(sub["Anio"].to_numpy(), sub["Casos"].to_numpy(), frac=smooth_frac)
            if sm is not None:
                sm_x, sm_y = sm
                fig.add_trace(
                    go.Scatter(
                        x=sm_x,
                        y=sm_y,
                        name=f"Suavizado (LOWESS) · {dept}",
                        mode="lines",
                        line=dict(width=2.5, color=color, shape="spline"),
                        hoverinfo="skip",
                    )
                )

    bp_summary = None
    if show_breakpoint and breakpoint_target:
        sub_bp = filtered_local[filtered_local["Departamento"] == breakpoint_target]
        bp = detect_breakpoint(sub_bp["Anio"].to_numpy(), sub_bp["Casos"].to_numpy())
        if bp is None:
            st.info(
                f"No hay suficientes años con datos en **{breakpoint_target}** "
                f"({site_val}) para estimar un punto de quiebre (se requieren al menos 6)."
            )
        else:
            bp_summary = bp
            fig.add_trace(
                go.Scatter(
                    x=bp.x_before, y=bp.y_pred_before, mode="lines",
                    line=dict(width=2, color="black", dash="dash"),
                    name=f"Tramo antes de {bp.year}", hoverinfo="skip",
                )
            )
            fig.add_trace(
                go.Scatter(
                    x=bp.x_after, y=bp.y_pred_after, mode="lines",
                    line=dict(width=2, color="black", dash="dash"),
                    name=f"Tramo desde {bp.year}", hoverinfo="skip",
                )
            )
            fig.add_vline(
                x=bp.year, line_width=1.5, line_dash="dash", line_color="black",
                annotation_text=f"Quiebre: {bp.year}", annotation_position="top",
            )

    mk_summary = None
    if show_mk and mk_target:
        sub_mk = dept_casos_series(mk_target, site_val, mk_year_range[0], mk_year_range[1])
        mkr = mann_kendall_trend(sub_mk["Anio"].to_numpy(), sub_mk["Casos"].to_numpy())
        if mkr is None:
            st.info(
                f"No hay suficientes años con datos en **{mk_target}** ({site_val}) "
                f"entre {mk_year_range[0]} y {mk_year_range[1]} para el test de "
                "Mann-Kendall (se requieren al menos 6)."
            )
        else:
            mk_summary = mkr
            fig.add_trace(
                go.Scatter(
                    x=mkr.x, y=mkr.y_sen, mode="lines",
                    line=dict(width=2, color="#0f9b8e", dash="dashdot"),
                    name=f"Pendiente de Sen · {mk_target}", hoverinfo="skip",
                )
            )

    arb_summary = None
    if show_arbitrary_break and arb_target and arb_break_year is not None:
        sub_arb = filtered_local[filtered_local["Departamento"] == arb_target]
        arb = chow_test_arbitrary_break(
            sub_arb["Anio"].to_numpy(),
            sub_arb["Casos"].to_numpy(),
            break_year=int(arb_break_year),
            implementation_lag=int(arb_lag),
        )
        if arb is None:
            st.info(
                f"No hay suficientes años antes/después de {int(arb_break_year)} "
                f"(considerando {int(arb_lag)} año(s) de implementación) en "
                f"**{arb_target}** ({site_val}) para aplicar el test (se requieren al "
                "menos 3 años a cada lado)."
            )
        else:
            arb_summary = arb
            fig.add_trace(
                go.Scatter(
                    x=arb.x_before, y=arb.y_pred_before, mode="lines",
                    line=dict(width=2, color="#e07a2c", dash="longdash"),
                    name=f"Antes del evento ({arb.break_year})", hoverinfo="skip",
                )
            )
            fig.add_trace(
                go.Scatter(
                    x=arb.x_after, y=arb.y_pred_after, mode="lines",
                    line=dict(width=2, color="#e07a2c", dash="longdash"),
                    name="Después del evento", hoverinfo="skip",
                )
            )
            fig.add_vline(
                x=arb.break_year, line_width=1.5, line_dash="dot", line_color="#e07a2c",
                annotation_text=f"Evento: {arb.break_year}", annotation_position="bottom",
            )
            if arb.implementation_lag:
                fig.add_vrect(
                    x0=arb.break_year, x1=arb.break_year + arb.implementation_lag,
                    fillcolor="#e07a2c", opacity=0.10, line_width=0,
                )

    fig.update_layout(
        title=f"Casos de cáncer — {site_val}",
        xaxis_title="Año",
        yaxis_title="N° de casos nuevos",
        yaxis_type="log" if log_scale else "linear",
        barmode="group",
        height=320 if compact else 580,
        template="plotly_white",
        legend=dict(
            orientation="h", yanchor="bottom", y=1.02, x=0,
            font=dict(size=10 if compact else 12),
        ),
        hovermode="x unified",
        margin=dict(t=60 if compact else 70, b=40),
        font=dict(size=11 if compact else 13),
    )
    all_years_local = sorted(filtered_local["Anio"].unique())
    fig.update_xaxes(
        tickmode="array", tickvals=all_years_local, tickangle=45,
        tickfont=dict(size=9 if compact else 11),
    )
    st.plotly_chart(fig, use_container_width=True, key=f"site_chart_{key_suffix}")

    if PERU_LABEL in depts_selected and peru_exclude:
        st.caption(
            f"ℹ️ El total de **{PERU_LABEL}** mostrado excluye: "
            f"{', '.join(peru_exclude)} (configurado en el panel lateral)."
        )

    summary_blocks = []
    if bp_summary is not None:
        sig_txt = (
            "**estadísticamente significativo**"
            if bp_summary.significant
            else "no alcanza significancia estadística"
        )
        summary_blocks.append(
            f"**📐 Punto de quiebre estimado para {breakpoint_target} ({site_val}): "
            f"año {bp_summary.year}** — el cambio de tendencia es {sig_txt} "
            f"(test de Chow: F = {bp_summary.f_stat:.2f}, p = {bp_summary.p_value:.4f}). "
            f"Pendiente antes: {bp_summary.slope_before:+.1f} casos/año · "
            f"pendiente después: {bp_summary.slope_after:+.1f} casos/año."
        )
    if mk_summary is not None:
        trend_es = {
            "increasing": "creciente 📈",
            "decreasing": "decreciente 📉",
            "no trend": "sin tendencia clara ➖",
        }.get(mk_summary.trend, mk_summary.trend)
        sig_txt = (
            "**estadísticamente significativa**"
            if mk_summary.significant
            else "no alcanza significancia estadística (α=0.05)"
        )
        summary_blocks.append(
            f"**🧪 Mann-Kendall para {mk_target} ({site_val}, "
            f"{mk_year_range[0]}–{mk_year_range[1]}): tendencia {trend_es}** — {sig_txt} "
            f"(z = {mk_summary.z_stat:.2f}, p = {mk_summary.p_value:.4f}, "
            f"τ de Kendall = {mk_summary.tau:.2f}). "
            f"Pendiente de Sen: {mk_summary.sen_slope:+.1f} casos/año "
            f"(n = {mk_summary.n_obs} años; robusta a outliers, a diferencia de la "
            f"pendiente de mínimos cuadrados). Método: {mk_summary.method}."
        )
    if arb_summary is not None:
        sig_txt = (
            "**estadísticamente significativo**"
            if arb_summary.significant
            else "no alcanza significancia estadística"
        )
        lag_txt = (
            f" (excluyendo {arb_summary.implementation_lag} año(s) de "
            "implementación tras el evento)"
            if arb_summary.implementation_lag
            else ""
        )
        summary_blocks.append(
            f"**🏛️ Quiebre por evento en {arb_target} ({site_val}): "
            f"año {arb_summary.break_year}**{lag_txt} — el cambio de tendencia es "
            f"{sig_txt} (test de Chow: F = {arb_summary.f_stat:.2f}, "
            f"p = {arb_summary.p_value:.4f}). "
            f"Pendiente antes: {arb_summary.slope_before:+.1f} casos/año · "
            f"pendiente después: {arb_summary.slope_after:+.1f} casos/año "
            f"(n antes = {arb_summary.n_before}, n después = {arb_summary.n_after})."
        )

    if summary_blocks:
        if compact:
            with st.expander("📊 Detalles del análisis estadístico"):
                for block in summary_blocks:
                    st.markdown(block)
        else:
            for block in summary_blocks:
                st.markdown(block)

    if not compact:
        st.markdown(
            '<p class="source-note">Nota: los valores corresponden a casos '
            "nuevos registrados por el INEN, no a tasas ajustadas por edad "
            "ni a la incidencia nacional total (el INEN es un centro de "
            "referencia, no cubre el 100% de los casos del país). "
            "Para tasas estandarizadas se requiere población por departamento y año.</p>",
            unsafe_allow_html=True,
        )


with tab_graph:
    render_site_analysis_chart(site, filtered, compact=False, key_suffix="main")

    # -----------------------------------------------------------------
    # Gráficos adicionales: comparar simultáneamente otras localizaciones
    # del tumor primario, cada una en su propio gráfico más compacto (2
    # por fila), con el mismo análisis estadístico y personalización de
    # colores que el gráfico principal.
    # -----------------------------------------------------------------
    extra_sites_to_plot = [s for s in extra_sites if s != site]
    if show_extra_sites and extra_sites_to_plot:
        st.markdown("---")
        st.markdown("#### Comparar con otras localizaciones")
        for row_start in range(0, len(extra_sites_to_plot), 2):
            row_sites = extra_sites_to_plot[row_start : row_start + 2]
            cols = st.columns(len(row_sites))
            for col, extra_site in zip(cols, row_sites):
                with col:
                    extra_frames = []
                    for dept in depts_selected:
                        sub_extra = dept_casos_series(
                            dept, extra_site, year_range[0], year_range[1]
                        ).copy()
                        sub_extra["Departamento"] = dept
                        extra_frames.append(sub_extra)
                    extra_filtered = (
                        pd.concat(extra_frames, ignore_index=True)
                        if extra_frames
                        else pd.DataFrame(columns=["Anio", "Casos", "Departamento"])
                    )
                    render_site_analysis_chart(
                        extra_site, extra_filtered, compact=True,
                        key_suffix=f"extra_{extra_site}",
                    )
with tab_table:
    pivot = filtered.pivot_table(
        index="Departamento", columns="Anio", values="Casos", aggfunc="sum"
    )
    st.dataframe(pivot.style.format("{:,.0f}", na_rep="—"), use_container_width=True)

with tab_ranking:
    st.markdown("#### Ranking de localizaciones de cáncer por año y departamento")
    st.caption(
        "Compara qué tipos de cáncer concentran más casos en un año y "
        "departamento específicos (similar al ranking por país de GLOBOCAN, "
        "aquí aplicado a localizaciones del tumor primario dentro de un departamento)."
    )
    years_all_desc = sorted(df["Anio"].unique().tolist(), reverse=True)
    depts_all = departamentos(df)
    all_sites = localizaciones(df)[1:]  # sin "Todas las localizaciones"

    rc1, rc2, rc3 = st.columns([1, 1.4, 1.2])
    with rc1:
        rank_year = st.selectbox("Año", options=years_all_desc, index=0, key="rank_year")
    with rc2:
        rank_dept = st.selectbox(
            "Departamento",
            options=depts_all,
            index=depts_all.index(PERU_LABEL),
            key="rank_dept",
        )
    with rc3:
        rank_top_n = st.slider("Top N localizaciones", 5, 38, 15, key="rank_top_n")

    rank_exclude = st.multiselect(
        "Excluir localización(es) del ranking",
        options=all_sites,
        default=[],
        key="rank_exclude",
        help=(
            "Útil para quitar categorías genéricas como 'Otros' o "
            "'Primario Desconocido', que suelen concentrar muchos casos "
            "sin ser clínicamente específicas, y así ver mejor el "
            "ranking de los tipos de cáncer concretos."
        ),
    )

    rank_data = get_year_ranking(rank_dept, rank_year, exclude=rank_exclude, top_n=rank_top_n)

    if rank_dept == PERU_LABEL and peru_exclude:
        st.caption(
            f"ℹ️ El total de Perú excluye: {', '.join(peru_exclude)} "
            "(configurado en el panel lateral)."
        )

    if rank_data.empty:
        st.warning(
            f"No hay datos registrados para **{rank_dept}** en **{rank_year}**"
            + (" con las localizaciones excluidas actuales." if rank_exclude else ".")
        )
    else:
        n_bars = len(rank_data)
        rank_base_color = (
            color_overrides.get(rank_dept, PRIMARY_COLOR) if custom_colors else PRIMARY_COLOR
        )
        bar_colors = shades_of(rank_base_color, n_bars)
        ranks = list(range(1, n_bars + 1))
        bar_text = [f"{r}° · {c:,.0f}" for r, c in zip(ranks, rank_data["Casos"])]

        fig_rank = go.Figure(
            go.Bar(
                x=rank_data["Casos"],
                y=rank_data["Localizacion"],
                orientation="h",
                marker_color=bar_colors,
                text=bar_text,
                textposition="outside",
                textfont=dict(size=15),
                hovertemplate="%{y}: %{x:,.0f} casos<extra></extra>",
            )
        )
        fig_rank.update_layout(
            title=f"Ranking de cánceres — {rank_dept}, {rank_year}",
            xaxis_title="N° de casos nuevos",
            yaxis_title="",
            yaxis=dict(autorange="reversed", tickfont=dict(size=13)),
            height=max(420, 26 * len(rank_data)),
            template="plotly_white",
            margin=dict(l=10, r=80, t=60, b=40),
        )
        st.plotly_chart(fig_rank, use_container_width=True)

        rank_table = rank_data[["Localizacion", "Casos"]].reset_index(drop=True)
        rank_table.insert(0, "Puesto", ranks)
        rank_table = rank_table.rename(
            columns={"Localizacion": "Localización del tumor primario"}
        )
        st.dataframe(
            rank_table.style.format({"Casos": "{:,.0f}"}),
            use_container_width=True,
            hide_index=True,
        )

        rank_csv = rank_table.to_csv(index=False).encode("utf-8")
        st.download_button(
            "Descargar ranking (CSV)",
            data=rank_csv,
            file_name=f"ranking_{rank_dept.replace(' ', '_')}_{rank_year}.csv",
            mime="text/csv",
        )

with tab_ranking_anim:
    st.markdown("#### Ranking animado — evolución de los cánceres en el tiempo")
    st.caption(
        "Muestra cómo cambia el ranking de localizaciones del tumor primario "
        "año a año ('bar chart race'). Reproduce la animación directamente "
        "aquí, o descárgala como GIF."
    )

    years_all_asc = sorted(df["Anio"].unique().tolist())

    ac1, ac2 = st.columns([1.6, 1])
    with ac1:
        anim_dept = st.selectbox(
            "Departamento",
            options=depts_all,
            index=depts_all.index(PERU_LABEL),
            key="anim_dept",
        )
    with ac2:
        anim_top_n = st.slider("Top N localizaciones", 5, 15, 8, key="anim_top_n")

    anim_year_range = st.select_slider(
        "Rango de años a animar",
        options=years_all_asc,
        value=(years_all_asc[0], years_all_asc[-1]),
        key="anim_year_range",
    )

    anim_exclude = st.multiselect(
        "Excluir localización(es) de la animación",
        options=all_sites,
        default=[],
        key="anim_exclude",
    )

    anim_dept_exclude: list[str] = []
    if anim_dept == PERU_LABEL:
        anim_dept_exclude = st.multiselect(
            "Excluir departamento(s) del total nacional (ej. Extranjero)",
            options=[d for d in depts_all if d != PERU_LABEL],
            default=peru_exclude,
            key="anim_dept_exclude",
            help=(
                "Igual que en el panel lateral, pero específico para esta "
                "animación. Por defecto trae lo configurado en el panel "
                "lateral; puedes cambiarlo aquí sin afectar el gráfico principal."
            ),
        )

    anim_seconds = st.slider(
        "Velocidad: segundos por año",
        min_value=0.2,
        max_value=2.0,
        value=0.7,
        step=0.1,
        key="anim_seconds",
        help="Valores bajos = animación más rápida. Aplica tanto a la vista interactiva como al GIF descargable.",
    )

    if anim_dept == PERU_LABEL and anim_dept_exclude:
        st.caption(f"ℹ️ El total de Perú excluye: {', '.join(anim_dept_exclude)}.")

    generate_anim = st.button("🎬 Generar ranking animado", key="anim_generate")

    if generate_anim:
        anim_years = list(range(anim_year_range[0], anim_year_range[1] + 1))
        if len(anim_years) < 2:
            st.warning("Selecciona un rango de al menos 2 años para animar.")
        else:
            with st.spinner("Generando animación..."):
                year_rankings = {}
                universe: set[str] = set()
                for y in anim_years:
                    r = get_year_ranking(
                        anim_dept, y, exclude=anim_exclude, peru_excl_override=anim_dept_exclude
                    )
                    year_rankings[y] = r.set_index("Localizacion")["Casos"]
                    universe.update(r.head(anim_top_n)["Localizacion"].tolist())
                universe = sorted(universe)

                if not universe:
                    st.warning(
                        f"No hay datos para **{anim_dept}** en el rango "
                        f"{anim_year_range[0]}–{anim_year_range[1]} con los "
                        "filtros actuales."
                    )
                else:
                    global_max = 0.0
                    for y in anim_years:
                        v = (
                            year_rankings[y]
                            .reindex(universe)
                            .fillna(0)
                            .sort_values(ascending=False)
                            .head(anim_top_n)
                        )
                        if len(v):
                            global_max = max(global_max, float(v.max()))
                    if global_max <= 0:
                        global_max = 1.0

                    # Duración de cada cuadro vs. duración de la transición:
                    # una transición que ocupa la mayor parte del intervalo
                    # (con easing) es lo que hace que el movimiento se vea
                    # fluido en vez de "saltar" de un año a otro.
                    frame_duration_ms = int(anim_seconds * 1000)
                    transition_ms = max(50, int(frame_duration_ms * 0.85))

                    year_annotation = dict(
                        text=str(anim_years[0]),
                        xref="x domain",
                        yref="y domain",
                        x=0.99,
                        y=0.03,
                        xanchor="right",
                        yanchor="bottom",
                        showarrow=False,
                        font=dict(size=58, color="black", family="Arial Black, Arial"),
                    )

                    anim_base_color = (
                        color_overrides.get(anim_dept, PRIMARY_COLOR)
                        if custom_colors
                        else PRIMARY_COLOR
                    )

                    frames = []
                    for y in anim_years:
                        vals = year_rankings[y].reindex(universe).fillna(0)
                        ordered = vals.sort_values(ascending=False).head(anim_top_n)
                        names = ordered.index.tolist()
                        values = ordered.values.tolist()
                        n = len(names)
                        y_pos = list(range(n, 0, -1))
                        colors = shades_of(anim_base_color, n)
                        text = [
                            f"{r}° {name}: {val:,.0f}"
                            for r, (name, val) in enumerate(zip(names, values), start=1)
                        ]
                        frame_annotation = dict(year_annotation, text=str(y))
                        frames.append(
                            go.Frame(
                                name=str(y),
                                data=[
                                    go.Bar(
                                        x=values,
                                        y=y_pos,
                                        orientation="h",
                                        marker_color=colors,
                                        text=text,
                                        textposition="outside",
                                        textfont=dict(size=14),
                                        hovertemplate="%{text}<extra></extra>",
                                    )
                                ],
                                layout=go.Layout(annotations=[frame_annotation]),
                            )
                        )

                    fig_anim = go.Figure(
                        data=frames[0].data,
                        layout=go.Layout(
                            title=f"Ranking de cánceres — {anim_dept}",
                            annotations=[year_annotation],
                            xaxis=dict(
                                title="N° de casos nuevos", range=[0, global_max * 1.2]
                            ),
                            yaxis=dict(
                                showticklabels=False,
                                range=[0.3, anim_top_n + 0.9],
                            ),
                            template="plotly_white",
                            height=max(420, 34 * anim_top_n) + 140,
                            margin=dict(l=20, r=160, t=70, b=170),
                            updatemenus=[
                                dict(
                                    type="buttons",
                                    direction="right",
                                    showactive=False,
                                    x=0.5,
                                    xanchor="center",
                                    y=-0.16,
                                    yanchor="top",
                                    pad=dict(t=10, r=10),
                                    buttons=[
                                        dict(
                                            label="▶ Reproducir",
                                            method="animate",
                                            args=[
                                                None,
                                                {
                                                    "frame": {
                                                        "duration": frame_duration_ms,
                                                        "redraw": True,
                                                    },
                                                    "fromcurrent": True,
                                                    "transition": {
                                                        "duration": transition_ms,
                                                        "easing": "cubic-in-out",
                                                    },
                                                },
                                            ],
                                        ),
                                        dict(
                                            label="⏸ Pausar",
                                            method="animate",
                                            args=[
                                                [None],
                                                {
                                                    "frame": {"duration": 0, "redraw": False},
                                                    "mode": "immediate",
                                                },
                                            ],
                                        ),
                                    ],
                                )
                            ],
                            sliders=[
                                dict(
                                    active=0,
                                    x=0.0,
                                    y=-0.28,
                                    len=1.0,
                                    currentvalue=dict(prefix="Año: ", font=dict(size=13)),
                                    steps=[
                                        dict(
                                            method="animate",
                                            label=str(y),
                                            args=[
                                                [str(y)],
                                                {
                                                    "frame": {
                                                        "duration": frame_duration_ms,
                                                        "redraw": True,
                                                    },
                                                    "mode": "immediate",
                                                    "transition": {
                                                        "duration": transition_ms,
                                                        "easing": "cubic-in-out",
                                                    },
                                                },
                                            ],
                                        )
                                        for y in anim_years
                                    ],
                                )
                            ],
                        ),
                        frames=frames,
                    )
                    st.plotly_chart(fig_anim, use_container_width=True)

                    if _GIF_EXPORT_AVAILABLE:
                        with st.spinner("Generando GIF descargable..."):
                            images = []
                            for y in anim_years:
                                vals = year_rankings[y].reindex(universe).fillna(0)
                                ordered = vals.sort_values(ascending=False).head(anim_top_n)
                                names = ordered.index.tolist()
                                values = ordered.values.tolist()
                                n = len(names)
                                colors = shades_of(anim_base_color, n)

                                mpl_fig, ax = plt.subplots(
                                    figsize=(9, 0.5 * anim_top_n + 1.4), dpi=110
                                )
                                ax.barh(range(n), values, color=colors)
                                ax.set_yticks(range(n))
                                ax.set_yticklabels(
                                    [f"{r}° {name}" for r, name in enumerate(names, start=1)],
                                    fontsize=12,
                                )
                                ax.set_xlim(0, global_max * 1.2)
                                for i, val in enumerate(values):
                                    ax.text(
                                        val + global_max * 0.01, i, f"{val:,.0f}",
                                        va="center", fontsize=11,
                                    )
                                ax.set_title(anim_dept, fontsize=13, loc="left", color="#555")
                                ax.text(
                                    0.98, 0.04, str(y),
                                    transform=ax.transAxes,
                                    fontsize=46, fontweight="black",
                                    color="black", ha="right", va="bottom",
                                )
                                ax.set_xlabel("N° de casos nuevos")
                                ax.invert_yaxis()
                                ax.spines[["top", "right"]].set_visible(False)
                                mpl_fig.tight_layout()

                                buf = io.BytesIO()
                                mpl_fig.savefig(buf, format="png")
                                plt.close(mpl_fig)
                                buf.seek(0)
                                images.append(imageio.imread(buf))

                            gif_buf = io.BytesIO()
                            imageio.mimsave(
                                gif_buf, images, format="GIF", duration=anim_seconds
                            )
                            gif_buf.seek(0)
                            gif_bytes = gif_buf.getvalue()

                        st.download_button(
                            "Descargar animación (GIF)",
                            data=gif_bytes,
                            file_name=f"ranking_animado_{anim_dept.replace(' ', '_')}_{anim_years[0]}_{anim_years[-1]}.gif",
                            mime="image/gif",
                        )
                    else:
                        st.info(
                            "La descarga en GIF no está disponible en este entorno "
                            "(faltan las librerías matplotlib/imageio), pero la "
                            "animación interactiva de arriba funciona igual."
                        )

def render_projection_chart(site_val: str, compact: bool = False, key_suffix: str = "") -> None:
    """Construye y renderiza la proyección de casos (3 escenarios +
    banda de incertidumbre) para una localización del tumor primario
    dada. La usan tanto la proyección principal como las proyecciones
    adicionales de otras localizaciones."""
    sub_proj = dept_casos_series(proj_target, site_val, year_range[0], year_range[1])
    proj = project_series(
        sub_proj["Anio"].to_numpy(), sub_proj["Casos"].to_numpy(), horizon=proj_horizon
    )
    if proj is None:
        st.warning(
            f"No hay suficientes años con datos en **{proj_target}** ({site_val}) "
            "para proyectar (se requieren al menos 6)."
        )
        return

    st.markdown(
        f"#### Proyección de casos — {site_val} · {proj_target} "
        f"({int(proj.years_future[0])}–{int(proj.years_future[-1])})"
    )
    if not compact:
        st.caption(proj.method_note)

    fig_proj = go.Figure()
    fig_proj.add_trace(
        go.Scatter(
            x=proj.years_hist, y=proj.values_hist, name="Histórico",
            mode="lines+markers", line=dict(width=2.0 if compact else 2.5, color=ACCENT_COLOR),
        )
    )
    fig_proj.add_trace(
        go.Scatter(
            x=np.concatenate([proj.years_future, proj.years_future[::-1]]),
            y=np.concatenate([proj.recommended_upper, proj.recommended_lower[::-1]]),
            fill="toself", fillcolor="rgba(47,111,168,0.15)",
            line=dict(color="rgba(0,0,0,0)"),
            name="Banda de incertidumbre (90%)", hoverinfo="skip",
        )
    )
    fig_proj.add_trace(
        go.Scatter(
            x=proj.years_future, y=proj.recommended,
            name="Recomendado (Holt amortiguado)",
            mode="lines+markers", line=dict(width=2.0 if compact else 2.5, color=SECONDARY_COLOR),
        )
    )
    fig_proj.add_trace(
        go.Scatter(
            x=proj.years_future, y=proj.conservative,
            name="Conservador (promedio reciente, plano)",
            mode="lines", line=dict(width=2, color="#7d7d7d", dash="dot"),
        )
    )
    fig_proj.add_trace(
        go.Scatter(
            x=proj.years_future, y=proj.optimistic,
            name="Lineal (pendiente de Sen, sin amortiguar)",
            mode="lines", line=dict(width=2, color="#c94141", dash="dash"),
        )
    )
    fig_proj.update_layout(
        xaxis_title="Año", yaxis_title="N° de casos (proyectado)",
        height=320 if compact else 520,
        template="plotly_white",
        legend=dict(
            orientation="h", yanchor="bottom", y=1.02, x=0,
            font=dict(size=10 if compact else 12),
        ),
        hovermode="x unified", margin=dict(t=40, b=40),
        font=dict(size=11 if compact else 13),
    )
    all_proj_years = list(proj.years_hist) + list(proj.years_future)
    fig_proj.update_xaxes(
        tickmode="array", tickvals=all_proj_years, tickangle=45,
        tickfont=dict(size=9 if compact else 11),
    )
    st.plotly_chart(fig_proj, use_container_width=True, key=f"proj_chart_{key_suffix}")

    table_proj = pd.DataFrame(
        {
            "Año": proj.years_future.astype(int),
            "Conservador": proj.conservative.round(0),
            "Recomendado": proj.recommended.round(0),
            "Recomendado (banda 90%)": [
                f"{lo:,.0f} – {hi:,.0f}"
                for lo, hi in zip(proj.recommended_lower, proj.recommended_upper)
            ],
            "Lineal (Sen)": proj.optimistic.round(0),
        }
    ).set_index("Año")

    if compact:
        with st.expander("📋 Tabla de escenarios"):
            st.dataframe(table_proj, use_container_width=True)
    else:
        st.markdown("##### Tabla de escenarios")
        st.dataframe(table_proj, use_container_width=True)

    proj_csv = table_proj.to_csv().encode("utf-8")
    st.download_button(
        "Descargar proyección (CSV)",
        data=proj_csv,
        file_name=f"proyeccion_{site_val.replace(' ', '_')}_{proj_target.replace(' ', '_')}.csv",
        mime="text/csv",
        key=f"proj_dl_{key_suffix}",
    )

    if not compact:
        st.markdown(
            '<p class="source-note">⚠️ Toda proyección es una extrapolación '
            "estadística del comportamiento histórico y no reemplaza la "
            "planificación epidemiológica basada en programas de tamizaje, "
            "cambios demográficos u otros factores no capturados por el "
            "modelo. Úsala como referencia de escenarios, no como cifra única.</p>",
            unsafe_allow_html=True,
        )


with tab_projection:
    if not show_projection or not proj_target:
        st.info(
            "Activa **'Proyectar casos a futuro'** en el panel lateral para ver "
            "esta sección. Se recomienda elegir una localización del tumor primario con serie "
            "relativamente completa (pocos años sin datos)."
        )
    else:
        render_projection_chart(site, compact=False, key_suffix="main")

        extra_sites_to_plot = [s for s in extra_sites if s != site]
        if show_extra_sites and extra_sites_to_plot:
            st.markdown("---")
            st.markdown("#### Proyección para otras localizaciones")
            for row_start in range(0, len(extra_sites_to_plot), 2):
                row_sites = extra_sites_to_plot[row_start : row_start + 2]
                cols = st.columns(len(row_sites))
                for col, extra_site in zip(cols, row_sites):
                    with col:
                        render_projection_chart(
                            extra_site, compact=True, key_suffix=f"extra_{extra_site}"
                        )

with tab_downloads:
    st.write("Descarga los datos filtrados actualmente en el panel:")
    csv_bytes = filtered.to_csv(index=False).encode("utf-8")
    st.download_button(
        "Descargar CSV (formato largo)",
        data=csv_bytes,
        file_name=f"cancer_inen_{site.replace(' ', '_')}.csv",
        mime="text/csv",
    )
    pivot_csv = pivot.to_csv().encode("utf-8")
    st.download_button(
        "Descargar CSV (tabla dinámica: departamento x año)",
        data=pivot_csv,
        file_name=f"cancer_inen_{site.replace(' ', '_')}_tabla.csv",
        mime="text/csv",
    )
    st.write("Descarga el dataset completo (todas las localizaciones del tumor primario y departamentos):")
    full_csv = df.to_csv(index=False).encode("utf-8")
    st.download_button(
        "Descargar dataset completo (CSV)",
        data=full_csv,
        file_name="cancer_inen_dataset_completo.csv",
        mime="text/csv",
    )

# ---------------------------------------------------------------------------
# Reporte ejecutivo (PDF / PPTX)
# ---------------------------------------------------------------------------


def _mpl_timeseries_chart(filtered_local: pd.DataFrame, depts: list[str], title: str) -> bytes:
    """Versión estática (matplotlib) del gráfico de casos por año y
    departamento, para incrustar en el reporte descargable."""
    fig, ax = plt.subplots(figsize=(9, 4.2), dpi=140)
    for i, dept in enumerate(depts):
        sub = filtered_local[filtered_local["Departamento"] == dept].sort_values("Anio")
        if sub["Casos"].dropna().empty:
            continue
        color = (
            color_overrides.get(dept, DEFAULT_PALETTE[i % len(DEFAULT_PALETTE)])
            if custom_colors
            else DEFAULT_PALETTE[i % len(DEFAULT_PALETTE)]
        )
        ax.plot(
            sub["Anio"], sub["Casos"], marker="o", markersize=3, linewidth=2,
            label=dept, color=color,
        )
    ax.set_title(title, fontsize=13, fontweight="bold", loc="left")
    ax.set_xlabel("Año")
    ax.set_ylabel("N° de casos nuevos")
    ax.legend(fontsize=8, loc="upper left", ncol=min(len(depts), 3))
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="y", alpha=0.3)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    plt.close(fig)
    buf.seek(0)
    return buf.getvalue()


def _mpl_rank_chart(rank_data_local: pd.DataFrame, title: str, base_color: str = PRIMARY_COLOR) -> bytes:
    """Versión estática (matplotlib) del ranking de localizaciones, para
    incrustar en el reporte descargable."""
    n = len(rank_data_local)
    colors = shades_of(base_color, n)
    fig, ax = plt.subplots(figsize=(9, max(3, 0.4 * n) + 1), dpi=140)
    ax.barh(range(n), rank_data_local["Casos"], color=colors)
    ax.set_yticks(range(n))
    ax.set_yticklabels(
        [f"{r}° {name}" for r, name in enumerate(rank_data_local["Localizacion"], start=1)],
        fontsize=9,
    )
    max_val = rank_data_local["Casos"].max()
    for i, val in enumerate(rank_data_local["Casos"]):
        ax.text(val + max_val * 0.01, i, f"{val:,.0f}", va="center", fontsize=8)
    ax.set_title(title, fontsize=13, fontweight="bold", loc="left")
    ax.set_xlabel("N° de casos nuevos")
    ax.invert_yaxis()
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    plt.close(fig)
    buf.seek(0)
    return buf.getvalue()


def _mpl_projection_chart(proj, title: str) -> bytes:
    """Versión estática (matplotlib) del gráfico de proyección, para
    incrustar en el reporte descargable."""
    fig, ax = plt.subplots(figsize=(9, 4.2), dpi=140)
    ax.plot(
        proj.years_hist, proj.values_hist, marker="o", markersize=3,
        linewidth=2, color=ACCENT_COLOR, label="Histórico",
    )
    ax.fill_between(
        proj.years_future, proj.recommended_lower, proj.recommended_upper,
        color=SECONDARY_COLOR, alpha=0.15, label="Banda 90%",
    )
    ax.plot(
        proj.years_future, proj.recommended, marker="o", markersize=3,
        linewidth=2, color=SECONDARY_COLOR, label="Recomendado",
    )
    ax.plot(
        proj.years_future, proj.conservative, linestyle=":", linewidth=1.6,
        color="#7d7d7d", label="Conservador",
    )
    ax.plot(
        proj.years_future, proj.optimistic, linestyle="--", linewidth=1.6,
        color="#c94141", label="Lineal (Sen)",
    )
    ax.set_title(title, fontsize=13, fontweight="bold", loc="left")
    ax.set_xlabel("Año")
    ax.set_ylabel("N° de casos (proyectado)")
    ax.legend(fontsize=8, loc="upper left")
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="y", alpha=0.3)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    plt.close(fig)
    buf.seek(0)
    return buf.getvalue()


def compute_site_summaries(site_val: str, filtered_local: pd.DataFrame) -> dict:
    """Calcula (sin renderizar nada en pantalla) los resúmenes
    estadísticos de una localización — tendencia/quiebre/Mann-Kendall/
    evento — según los controles del panel lateral, para el reporte."""
    bp_summary = None
    if show_breakpoint and breakpoint_target:
        sub_bp = filtered_local[filtered_local["Departamento"] == breakpoint_target]
        bp_summary = detect_breakpoint(sub_bp["Anio"].to_numpy(), sub_bp["Casos"].to_numpy())

    mk_summary = None
    if show_mk and mk_target:
        sub_mk = dept_casos_series(mk_target, site_val, mk_year_range[0], mk_year_range[1])
        mk_summary = mann_kendall_trend(sub_mk["Anio"].to_numpy(), sub_mk["Casos"].to_numpy())

    arb_summary = None
    if show_arbitrary_break and arb_target and arb_break_year is not None:
        sub_arb = filtered_local[filtered_local["Departamento"] == arb_target]
        arb_summary = chow_test_arbitrary_break(
            sub_arb["Anio"].to_numpy(), sub_arb["Casos"].to_numpy(),
            break_year=int(arb_break_year), implementation_lag=int(arb_lag),
        )
    return {"bp": bp_summary, "mk": mk_summary, "arb": arb_summary}


def format_summary_lines(site_val: str, summaries: dict) -> list[str]:
    """Texto plano de los resúmenes estadísticos, para el reporte."""
    lines = []
    bp = summaries.get("bp")
    if bp is not None:
        sig = "estadísticamente significativo" if bp.significant else "no alcanza significancia estadística"
        lines.append(
            f"Quiebre estimado para {breakpoint_target} ({site_val}): año {bp.year} — {sig} "
            f"(test de Chow: F={bp.f_stat:.2f}, p={bp.p_value:.4f}). "
            f"Pendiente antes: {bp.slope_before:+.1f} casos/año, después: {bp.slope_after:+.1f} casos/año."
        )
    mk = summaries.get("mk")
    if mk is not None:
        trend_es = {
            "increasing": "creciente", "decreasing": "decreciente", "no trend": "sin tendencia clara",
        }.get(mk.trend, mk.trend)
        sig = "significativa" if mk.significant else "no significativa"
        lines.append(
            f"Mann-Kendall para {mk_target} ({site_val}, {mk_year_range[0]}–{mk_year_range[1]}): "
            f"tendencia {trend_es} — {sig} (p={mk.p_value:.4f}). "
            f"Pendiente de Sen: {mk.sen_slope:+.1f} casos/año."
        )
    arb = summaries.get("arb")
    if arb is not None:
        sig = "estadísticamente significativo" if arb.significant else "no alcanza significancia estadística"
        lines.append(
            f"Quiebre por evento en {arb_target} ({site_val}): año {arb.break_year} — {sig} "
            f"(test de Chow: F={arb.f_stat:.2f}, p={arb.p_value:.4f})."
        )
    return lines


def build_report_sections(
    inc_kpis: bool, inc_main_chart: bool, inc_extra_charts: bool, inc_stats: bool,
    inc_ranking: bool, inc_projection: bool, inc_table: bool,
) -> list[dict]:
    """Arma la lista de secciones del reporte (cada una con encabezado,
    texto, imagen y/o tabla) a partir de las casillas elegidas por el
    usuario, reutilizando exactamente los mismos datos y ajustes
    (localización, departamentos, exclusiones, análisis, proyección)
    configurados en el resto del dashboard."""
    sections: list[dict] = []

    if inc_kpis:
        kpi_lines = [
            f"Localización del tumor primario: {site}",
            f"Departamento(s) de residencia: {format_region_list(depts_selected, max_show=10)}",
            f"Variación {year_a} → {year_b}: " + (f"{delta_pct:+.1f}%" if pd.notna(delta_pct) else "s/d"),
            "Cambio absoluto: " + (f"{delta_abs:+,.0f} casos" if pd.notna(delta_abs) else "s/d"),
            "CAGR (crecimiento anual compuesto): " + (f"{cagr * 100:+.1f}%/año" if pd.notna(cagr) else "s/d"),
            "Año pico: " + (f"{peak_year} · {peak_val:,.0f} casos" if peak_year is not None else "s/d"),
        ]
        if mk_kpi is not None:
            trend_word = {
                "increasing": "Creciente", "decreasing": "Decreciente", "no trend": "Sin tendencia clara",
            }.get(mk_kpi.trend, mk_kpi.trend)
            kpi_lines.append(f"Tendencia general (Mann-Kendall): {trend_word} (p={mk_kpi.p_value:.3f})")
        if PERU_LABEL in depts_selected and peru_exclude:
            kpi_lines.append(f"El total de {PERU_LABEL} excluye: {', '.join(peru_exclude)}.")
        sections.append({"heading": "Resumen ejecutivo", "text": kpi_lines})

    if inc_main_chart:
        img = _mpl_timeseries_chart(filtered, depts_selected, f"Casos de cáncer — {site}")
        sec = {"heading": f"Gráfico — {site}", "image_bytes": img}
        if inc_stats:
            lines = format_summary_lines(site, compute_site_summaries(site, filtered))
            if lines:
                sec["text"] = lines
        sections.append(sec)

    if inc_extra_charts and show_extra_sites:
        for extra_site in [s for s in extra_sites if s != site]:
            extra_frames = []
            for dept in depts_selected:
                sub_extra = dept_casos_series(dept, extra_site, year_range[0], year_range[1]).copy()
                sub_extra["Departamento"] = dept
                extra_frames.append(sub_extra)
            extra_filtered = (
                pd.concat(extra_frames, ignore_index=True) if extra_frames else pd.DataFrame()
            )
            if extra_filtered.empty or extra_filtered["Casos"].dropna().empty:
                continue
            img = _mpl_timeseries_chart(extra_filtered, depts_selected, f"Casos de cáncer — {extra_site}")
            sec = {"heading": f"Gráfico — {extra_site}", "image_bytes": img}
            if inc_stats:
                lines = format_summary_lines(extra_site, compute_site_summaries(extra_site, extra_filtered))
                if lines:
                    sec["text"] = lines
            sections.append(sec)

    if inc_ranking:
        rank_data_report = get_year_ranking(rank_dept, rank_year, exclude=rank_exclude, top_n=rank_top_n)
        if not rank_data_report.empty:
            base_color = color_overrides.get(rank_dept, PRIMARY_COLOR) if custom_colors else PRIMARY_COLOR
            img = _mpl_rank_chart(
                rank_data_report, f"Ranking de cánceres — {rank_dept}, {rank_year}", base_color=base_color
            )
            table_data = [["Puesto", "Localización", "Casos"]] + [
                [str(i + 1), row.Localizacion, f"{row.Casos:,.0f}"]
                for i, row in enumerate(rank_data_report.itertuples())
            ]
            sections.append({
                "heading": f"Ranking — {rank_dept}, {rank_year}",
                "image_bytes": img,
                "table_data": table_data,
            })

    if inc_projection and show_projection and proj_target:
        sub_proj = dept_casos_series(proj_target, site, year_range[0], year_range[1])
        proj = project_series(sub_proj["Anio"].to_numpy(), sub_proj["Casos"].to_numpy(), horizon=proj_horizon)
        if proj is not None:
            img = _mpl_projection_chart(proj, f"Proyección — {site} · {proj_target}")
            table_data = [["Año", "Conservador", "Recomendado", "Banda 90%", "Lineal (Sen)"]] + [
                [str(int(y)), f"{c:,.0f}", f"{r:,.0f}", f"{lo:,.0f}–{hi:,.0f}", f"{o:,.0f}"]
                for y, c, r, lo, hi, o in zip(
                    proj.years_future, proj.conservative, proj.recommended,
                    proj.recommended_lower, proj.recommended_upper, proj.optimistic,
                )
            ]
            sections.append({
                "heading": f"Proyección — {site} · {proj_target}",
                "image_bytes": img,
                "table_data": table_data,
                "text": [proj.method_note],
            })

    if inc_table:
        # Se transpone (años como filas) para que la tabla quepa en el
        # ancho de la página A4 y se pagine sola si hay muchos años —
        # con años como columnas, una serie larga se corta al borde.
        pivot_report = filtered.pivot_table(index="Departamento", columns="Anio", values="Casos", aggfunc="sum")
        header = ["Año"] + list(pivot_report.index)
        rows = [header]
        for year_col in pivot_report.columns:
            row_vals = [
                f"{v:,.0f}" if pd.notna(v) else "—" for v in pivot_report[year_col]
            ]
            rows.append([str(int(year_col))] + row_vals)
        sections.append({"heading": f"Tabla de datos — {site}", "table_data": rows})

    return sections


def build_pdf_report(sections: list[dict], report_title: str) -> bytes:
    """Arma el PDF (A4) del reporte con reportlab (sin dependencias de
    sistema como Chrome/LibreOffice). Cada gráfico o tabla ocupa su
    propia página — si una sección tiene ambos (p. ej. Proyección), el
    gráfico va en una página y la tabla en la siguiente."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
        leftMargin=1.8 * cm, rightMargin=1.8 * cm,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleCustom", parent=styles["Title"], textColor=rl_colors.HexColor(ACCENT_COLOR),
    )
    h2_style = ParagraphStyle(
        "H2Custom", parent=styles["Heading2"], textColor=rl_colors.HexColor(ACCENT_COLOR),
        spaceBefore=14,
    )
    body_style = styles["BodyText"]

    story = [
        Paragraph(report_title, title_style),
        Paragraph(f"Generado el {datetime.now():%d/%m/%Y %H:%M}", styles["Normal"]),
    ]

    for sec in sections:
        story.append(PageBreak())
        story.append(Paragraph(sec["heading"], h2_style))
        if sec.get("text"):
            for line in sec["text"]:
                story.append(Paragraph(line, body_style))
            story.append(Spacer(1, 0.3 * cm))

        if sec.get("image_bytes"):
            # El gráfico usa la mayor parte de la página, ya que tiene
            # toda la hoja para él solo.
            story.append(RLImage(io.BytesIO(sec["image_bytes"]), width=17 * cm, height=9.6 * cm))

        if sec.get("table_data"):
            if sec.get("image_bytes"):
                # El gráfico ya ocupó esta página: la tabla pasa a la
                # siguiente, cada elemento en su propia hoja.
                story.append(PageBreak())
                story.append(Paragraph(f"{sec['heading']} — tabla", h2_style))
            t = Table(sec["table_data"], hAlign="LEFT", repeatRows=1)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), rl_colors.HexColor(ACCENT_COLOR)),
                ("TEXTCOLOR", (0, 0), (-1, 0), rl_colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("GRID", (0, 0), (-1, -1), 0.4, rl_colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [rl_colors.white, rl_colors.HexColor("#f7f7fb")]),
            ]))
            story.append(t)

    doc.build(story)
    buf.seek(0)
    return buf.getvalue()


def build_pptx_report(sections: list[dict], report_title: str) -> bytes:
    """Arma la presentación PPTX del reporte con python-pptx (sin
    dependencias de sistema). Cada gráfico o tabla ocupa su propia
    diapositiva — si una sección tiene ambos (p. ej. Proyección), el
    gráfico va en una diapositiva y la tabla en la siguiente."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    accent_rgb = RGBColor.from_string(ACCENT_COLOR.lstrip("#"))

    def add_heading(slide, heading_text: str) -> None:
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.8))
        title_box.text_frame.text = heading_text
        title_box.text_frame.paragraphs[0].font.size = Pt(28)
        title_box.text_frame.paragraphs[0].font.bold = True
        title_box.text_frame.paragraphs[0].font.color.rgb = accent_rgb

    def add_text_block(slide, lines: list[str] | None, y_cursor: float) -> float:
        if not lines:
            return y_cursor
        body_box = slide.shapes.add_textbox(Inches(0.5), Inches(y_cursor), Inches(12.3), Inches(1.8))
        tf = body_box.text_frame
        tf.word_wrap = True
        # Estima cuántas líneas ocupará cada párrafo al envolverse, para
        # reservar el espacio real (una nota larga puede ocupar 3-4
        # líneas) y que el gráfico de abajo no quede superpuesto.
        chars_per_line = 128  # aprox. para Pt(14) en un ancho de 12.3"
        total_lines = 0
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.size = Pt(14)
            total_lines += max(1, -(-len(line) // chars_per_line))  # ceil
        return y_cursor + 0.30 + 0.27 * total_lines

    def add_table(slide, table_data: list[list[str]], y_cursor: float) -> None:
        # Las diapositivas no paginan: si la tabla es muy larga, se
        # muestran solo las filas más recientes con una nota (el PDF sí
        # incluye la tabla completa, paginada automáticamente).
        max_rows_per_slide = 16
        full_len = len(table_data)
        if full_len > max_rows_per_slide:
            header_row = table_data[0]
            kept_rows = table_data[-(max_rows_per_slide - 1):]
            table_data = [header_row] + kept_rows
            note_box = slide.shapes.add_textbox(Inches(0.7), Inches(y_cursor), Inches(11.9), Inches(0.35))
            note_box.text_frame.text = (
                f"Mostrando los últimos {len(kept_rows)} registros de "
                f"{full_len - 1} (tabla completa en la versión PDF)."
            )
            note_box.text_frame.paragraphs[0].font.size = Pt(11)
            note_box.text_frame.paragraphs[0].font.italic = True
            y_cursor += 0.4

        rows = len(table_data)
        cols = len(table_data[0])
        table_shape = slide.shapes.add_table(
            rows, cols, Inches(0.7), Inches(y_cursor), Inches(11.9), Inches(min(5.5, 0.35 * rows))
        )
        table = table_shape.table
        for r, row_vals in enumerate(table_data):
            for c, val in enumerate(row_vals):
                cell = table.cell(r, c)
                cell.text = str(val)
                for para in cell.text_frame.paragraphs:
                    para.font.size = Pt(10)

    slide = prs.slides.add_slide(blank)
    tx = slide.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(11.7), Inches(1.5))
    tf = tx.text_frame
    tf.text = report_title
    tf.paragraphs[0].font.size = Pt(40)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = accent_rgb
    sub = slide.shapes.add_textbox(Inches(0.8), Inches(4.0), Inches(11.7), Inches(0.8))
    sub.text_frame.text = f"Generado el {datetime.now():%d/%m/%Y %H:%M}"

    for sec in sections:
        slide = prs.slides.add_slide(blank)
        add_heading(slide, sec["heading"])
        y_cursor = add_text_block(slide, sec.get("text"), 1.2)

        if sec.get("image_bytes"):
            # Se fija solo la altura (no el ancho) para que la imagen
            # siempre quepa en el espacio restante de la diapositiva sin
            # superponerse al texto de arriba, usando la mayor parte de
            # la diapositiva ya que tiene toda la lámina para ella sola.
            available_height = max(2.0, 7.5 - y_cursor - 0.3)
            slide.shapes.add_picture(
                io.BytesIO(sec["image_bytes"]), Inches(0.7), Inches(y_cursor),
                height=Inches(min(5.8, available_height)),
            )

        if sec.get("table_data"):
            if sec.get("image_bytes"):
                # El gráfico ya ocupó esta diapositiva: la tabla pasa a
                # una nueva, cada elemento en su propia lámina.
                slide = prs.slides.add_slide(blank)
                add_heading(slide, f"{sec['heading']} — tabla")
                y_cursor = 1.2
            add_table(slide, sec["table_data"], y_cursor)

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf.getvalue()


with tab_report:
    st.markdown("#### Reporte ejecutivo descargable")
    st.caption(
        "Arma un reporte con las secciones que elijas y descárgalo en PDF "
        "(A4) o PowerPoint, listo para compartir o presentar."
    )

    if not (_GIF_EXPORT_AVAILABLE and (_PDF_EXPORT_AVAILABLE or _PPTX_EXPORT_AVAILABLE)):
        st.warning(
            "La generación de reportes requiere las librerías matplotlib, "
            "reportlab y python-pptx, que no están disponibles en este "
            "entorno. Instálalas con `pip install -r requirements.txt` y "
            "reinicia la app."
        )
    else:
        extra_sites_to_plot_report = [s for s in extra_sites if s != site]

        st.markdown("##### Elige qué incluir")
        rc1, rc2 = st.columns(2)
        with rc1:
            inc_kpis = st.checkbox("Portada con resumen (KPIs)", value=True, key="inc_kpis")
            inc_main_chart = st.checkbox(
                f"Gráfico principal — {site}", value=True, key="inc_main_chart"
            )
            inc_extra_charts = st.checkbox(
                "Gráficos adicionales de otras localizaciones",
                value=bool(show_extra_sites and extra_sites_to_plot_report),
                disabled=not (show_extra_sites and extra_sites_to_plot_report),
                key="inc_extra_charts",
            )
            inc_stats = st.checkbox(
                "Análisis estadístico (tendencia, quiebre, Mann-Kendall, evento)",
                value=any([show_breakpoint, show_mk, show_arbitrary_break]),
                key="inc_stats",
            )
        with rc2:
            inc_ranking = st.checkbox(
                f"Ranking por año — {rank_dept}, {rank_year}", value=True, key="inc_ranking"
            )
            inc_projection = st.checkbox(
                "Proyección de casos",
                value=bool(show_projection and proj_target),
                disabled=not (show_projection and proj_target),
                key="inc_projection",
            )
            inc_table = st.checkbox(
                "Tabla de datos (departamento × año)", value=False, key="inc_table"
            )

        report_title = st.text_input(
            "Título del reporte", value=f"Cáncer en el Tiempo — {site}", key="report_title"
        )

        gen_col1, gen_col2 = st.columns(2)
        with gen_col1:
            gen_pdf = st.button(
                "📄 Generar PDF (A4)", key="gen_pdf",
                disabled=not _PDF_EXPORT_AVAILABLE, use_container_width=True,
            )
        with gen_col2:
            gen_pptx = st.button(
                "📊 Generar PPTX", key="gen_pptx",
                disabled=not _PPTX_EXPORT_AVAILABLE, use_container_width=True,
            )

        if gen_pdf or gen_pptx:
            if not any([inc_kpis, inc_main_chart, inc_extra_charts, inc_ranking, inc_projection, inc_table]):
                st.warning("Selecciona al menos una sección para incluir en el reporte.")
            else:
                with st.spinner("Generando reporte..."):
                    sections = build_report_sections(
                        inc_kpis, inc_main_chart, inc_extra_charts, inc_stats,
                        inc_ranking, inc_projection, inc_table,
                    )
                if not sections:
                    st.warning("No hay datos suficientes para generar el reporte con la selección actual.")
                else:
                    if gen_pdf:
                        pdf_bytes = build_pdf_report(sections, report_title)
                        st.download_button(
                            "⬇️ Descargar PDF", data=pdf_bytes,
                            file_name=f"{report_title.replace(' ', '_')}.pdf",
                            mime="application/pdf", key="dl_pdf",
                        )
                    if gen_pptx:
                        pptx_bytes = build_pptx_report(sections, report_title)
                        st.download_button(
                            "⬇️ Descargar PPTX", data=pptx_bytes,
                            file_name=f"{report_title.replace(' ', '_')}.pptx",
                            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                            key="dl_pptx",
                        )

# ---------------------------------------------------------------------------
# Autoría
# ---------------------------------------------------------------------------

st.divider()
st.markdown(
    """
    <p style="text-align:center; color:#666; font-size:0.85rem;">
    © Luis A. Orrego Ferreyros, DDS, Econ., MCE, MMD, PhD(c), CQRM ·
    Epidemiólogo y Economista de la Salud · INEN
    </p>
    """,
    unsafe_allow_html=True,
)
